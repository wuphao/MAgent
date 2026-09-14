from pathlib import Path
import re,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
base=Path(__file__).parent
refs=[
'XUE C, KOWSHIK S S, LTEIF D, et al. AI-based differential diagnosis of dementia etiologies on multimodal data[J]. Nature Medicine, 2024, 30(10): 2977-2989. DOI:10.1038/s41591-024-03118-z.',
'YANG Y, YUAN Y, ZHANG G, et al. Artificial intelligence-enabled detection and assessment of Parkinson’s disease using nocturnal breathing signals[J]. Nature Medicine, 2022, 28(10): 2207-2215. DOI:10.1038/s41591-022-01932-x.',
'SCHALKAMP A K, PEALL K J, HARRISON N A, et al. Wearable movement-tracking data identify Parkinson’s disease years before clinical diagnosis[J]. Nature Medicine, 2023, 29: 2048-2056. DOI:10.1038/s41591-023-02440-2.',
'AZADMALEKI H, HAGHBIN Y, RASHIDI S, et al. SpeechCARE: dynamic multimodal modeling for cognitive screening in diverse linguistic and speech task contexts[J]. npj Digital Medicine, 2025, 8(1): 677. DOI:10.1038/s41746-025-02026-x.',
'YADA Y, NAOKI H. Decomposing heterogeneity in disease progression speeds and pathways[J]. npj Digital Medicine, 2026, 9(1): 562. DOI:10.1038/s41746-026-02665-8.',
'OEHRN C R, CERNERA S, HAMMER L H, et al. Chronic adaptive deep brain stimulation versus conventional stimulation in Parkinson’s disease: a blinded randomized feasibility trial[J]. Nature Medicine, 2024, 30(11): 3345-3356. DOI:10.1038/s41591-024-03196-z.',
'SINGHAL K, AZIZI S, TU T, et al. Large language models encode clinical knowledge[J]. Nature, 2023, 620(7972): 172-180. DOI:10.1038/s41586-023-06291-2.',
'CRAWFORD J L. Linguistic changes in spontaneous speech for detecting Parkinson’s disease using large language models[J]. PLOS Digital Health, 2025, 4(2): e0000757. DOI:10.1371/journal.pdig.0000757.',
'CASTELLI M, SOUSA M, VOJTECH I, et al. Detecting neuropsychiatric fluctuations in Parkinson’s Disease using patients’ own words: the potential of large language models[J]. npj Parkinson’s Disease, 2025, 11(1): 79. DOI:10.1038/s41531-025-00939-8.',
'TIAN J, FARD P, CAGAN C, et al. An autonomous agentic workflow for clinical detection of cognitive concerns using large language models[J]. npj Digital Medicine, 2026, 9(1): 51. DOI:10.1038/s41746-025-02324-4.',
'LAHIRI A K, HU Q V. AlzheimerRAG: multimodal retrieval-augmented generation for clinical use cases[J]. Machine Learning and Knowledge Extraction, 2025, 7(3): 89. DOI:10.3390/make7030089.',
'KOWSHIK S S, JASODANAND V H, BELLITTI M, et al. Domain-adapted language model using reinforcement learning for various dementias[PP/OL]. medRxiv (2026-03-23)[2026-09-14]. https://www.medrxiv.org/content/10.64898/2026.03.17.26348154v1. DOI:10.64898/2026.03.17.26348154.',
'HARRISON J R, ROBERTSON A, TANG S L, et al. Automating collateral histories in dementia: development and proof-of-concept evaluation of the LUMEN conversational AI[J/OL]. International Psychogeriatrics, 2026: 100221[2026-09-14]. https://doi.org/10.1016/j.inpsyc.2026.100221. DOI:10.1016/j.inpsyc.2026.100221.',
'HAGER P, JUNGMANN F, HOLLAND R, et al. Evaluation and mitigation of the limitations of large language models in clinical decision-making[J]. Nature Medicine, 2024, 30(9): 2613-2622. DOI:10.1038/s41591-024-03097-1.'
]
doc=Document();sec=doc.sections[0]
for element in doc.styles.element.xpath('.//w:pBdr'):
 element.getparent().remove(element)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.5);sec.left_margin=Cm(2.8);sec.right_margin=Cm(2.5)
sec.footer_distance=Cm(1.3)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
 s.font.size=Pt(12)
 s.paragraph_format.widow_control=True
n=doc.styles['Normal'];n.paragraph_format.line_spacing=1.5;n.paragraph_format.space_after=Pt(0);n.paragraph_format.first_line_indent=Pt(24)
for name,size in [('Title',18),('Heading 1',15),('Heading 2',13)]:
 s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.first_line_indent=Pt(0);s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
doc.styles['Title'].paragraph_format.space_before=Pt(0)
def runs(p,text,super_cite=True):
 for part in re.split(r'(\[\d+(?:,\d+)*\])',text):
  r=p.add_run(part)
  if super_cite and re.fullmatch(r'\[\d+(?:,\d+)*\]',part):r.font.superscript=True;r.font.size=Pt(9)
for line in base.joinpath('content.txt').read_text(encoding='utf-8').splitlines():
 if line.startswith('# '):
  p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
 elif line.strip():
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;runs(p,line)
p=doc.add_paragraph('参考文献','Heading 1');p.paragraph_format.page_break_before=True
for i,ref in enumerate(refs,1):
 p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(-24);p.paragraph_format.left_indent=Pt(24);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_together=True
 r=p.add_run(f'[{i}] {ref}');r.font.size=Pt(10.5)
 # Bookmark each entry so inline citations have stable targets for later field updates.
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(i));start.set(qn('w:name'),f'ref_{i}');p._p.insert(0,start)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(i));p._p.append(end)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
doc.core_properties.title='神经退行性疾病人工智能辅助诊疗研究现状'
doc.core_properties.subject='人工智能与大语言模型辅助诊疗研究综述'
doc.core_properties.author='';doc.core_properties.keywords='神经退行性疾病;人工智能;大语言模型;辅助诊断;检索增强生成'
out=base.parent/'神经退行性疾病人工智能辅助诊疗研究现状.docx';doc.save(out)
text=base.joinpath('content.txt').read_text(encoding='utf-8');seen=[]
for group in re.findall(r'\[(\d+(?:,\d+)*)\]',text):
 for num in map(int,group.split(',')):
  if num not in seen:seen.append(num)
assert seen==list(range(1,len(refs)+1)),seen
base.joinpath('reference_audit.json').write_text(json.dumps({'reference_count':len(refs),'first_appearance_order':seen,'standard':'GB/T 7714-2025 sequential numeric','preprints':[12],'online_first':[13],'body_chinese_characters':len(re.findall('[\u4e00-\u9fff]',text))},ensure_ascii=False,indent=2),encoding='utf-8')
print('DOCX created; 14 references; citation sequence verified')
