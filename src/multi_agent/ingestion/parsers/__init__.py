from __future__ import annotations

import csv
import importlib.util
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from multi_agent.domain.mapping import ParsedDocument, ParsedRecord
from multi_agent.ingestion.manifest import InputManifest
from multi_agent.storage.assets import AssetManifest, AssetRepository
from multi_agent.ingestion.parsers.text import TextDocument, TextParser, TextSpan


class Parser:
    version = "stage02-parser/1"

    def __init__(self, assets: AssetRepository):
        self.assets = assets

    def parse(self, manifest: InputManifest) -> ParsedDocument:
        path = manifest.path
        media_type = manifest.media_type or _media_type(path)
        asset = self.assets.register(
            AssetManifest(
                project_id=manifest.project_id,
                path=manifest.path,
                source_namespace=manifest.source_namespace,
                media_type=media_type,
            )
        )
        if media_type == "application/json":
            root = json.loads(path.read_text(encoding="utf-8"))
            records = [ParsedRecord(data=root, record_locator="$", value_locators=_json_value_locators(root))]
            return ParsedDocument(
                project_id=manifest.project_id,
                source_namespace=manifest.source_namespace,
                asset_id=asset.asset_id,
                asset_revision=asset.revision,
                media_type=media_type,
                available_at=asset.available_at,
                root=root,
                records=records,
                parser_version=self.version,
                layout="json",
            )
        if media_type == "text/csv":
            return self._parse_csv(path, manifest, asset.asset_id, asset.revision, media_type)
        if media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return self._parse_xlsx(path, manifest, asset.asset_id, asset.revision, media_type)
        raise ValueError(f"unsupported media_type: {media_type}")

    def _parse_csv(self, path: Path, manifest: InputManifest, asset_id: str, asset_revision: int, media_type: str) -> ParsedDocument:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames:
                raise ValueError("csv has no header")
            rows = list(reader)
        if not rows:
            raise ValueError("csv has no data rows")
        records = []
        for index, row in enumerate(rows, start=2):
            locators = {field: f"row:{index}:col:{field}" for field in row}
            records.append(ParsedRecord(data=row, record_locator=f"row:{index}", value_locators=locators))
        return ParsedDocument(
            project_id=manifest.project_id,
            source_namespace=manifest.source_namespace,
            asset_id=asset_id,
            asset_revision=asset_revision,
            media_type=media_type,
            available_at=self.assets.get(manifest.project_id, asset_id, asset_revision).available_at,
            root={"rows": rows},
            records=records,
            parser_version=self.version,
            layout="csv",
        )

    def _parse_xlsx(self, path: Path, manifest: InputManifest, asset_id: str, asset_revision: int, media_type: str) -> ParsedDocument:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("xlsx has no rows")
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        if not any(headers):
            raise ValueError("xlsx has no header")
        if len(headers) != len(set(headers)):
            raise ValueError("xlsx has duplicate headers")
        records = []
        raw_rows = []
        for row_number, values in enumerate(rows[1:], start=2):
            row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
            raw_rows.append(row)
            locators = {field: f"sheet:{sheet.title}:row:{row_number}:col:{field}" for field in row}
            records.append(ParsedRecord(data=row, record_locator=f"sheet:{sheet.title}:row:{row_number}", value_locators=locators))
        if not records:
            raise ValueError("xlsx has no data rows")
        return ParsedDocument(
            project_id=manifest.project_id,
            source_namespace=manifest.source_namespace,
            asset_id=asset_id,
            asset_revision=asset_revision,
            media_type=media_type,
            available_at=self.assets.get(manifest.project_id, asset_id, asset_revision).available_at,
            root={"sheet": sheet.title, "rows": raw_rows},
            records=records,
            parser_version=self.version,
            layout="xlsx",
        )


def _media_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".json":
        return "application/json"
    if suffix == ".csv":
        return "text/csv"
    if suffix == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    raise ValueError(f"cannot infer media_type from {path}")


def _json_value_locators(value: Any, prefix: str = "$") -> dict[str, str]:
    locators: dict[str, str] = {}
    if isinstance(value, dict):
        for key, item in value.items():
            child = f"{prefix}.{key}"
            locators[child] = child
            locators.update(_json_value_locators(item, child))
    elif isinstance(value, list):
        for index, item in enumerate(value):
            child = f"{prefix}[{index}]"
            locators[child] = child
            locators.update(_json_value_locators(item, child))
    return locators


__all__ = ["Parser", "TextDocument", "TextParser", "TextSpan"]
